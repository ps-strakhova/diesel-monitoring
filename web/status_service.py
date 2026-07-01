"""
Сервис для работы со статусом
"""

import sqlite3
from typing import Optional, List, Dict, Any
from datetime import datetime, timedelta
import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from .database import get_controller_display_name
import logging

logger = logging.getLogger(__name__)

class StatusService:
    def __init__(self, db_path):
        self.db_path = db_path
        self.snapshot_params = [
            (1000, "Обороты", "об/мин"),
            (1036, "АКБ", "В"),
            (1038, "Масло", "Бар"),
            (1039, "ОЖ", "°C"),
            (1040, "Топливо", "%"),
            (1017, "G-Частота", "Гц"),
        ]
        self.full_status_params = {
            'rpm': 1000, 'p_load': 1001, 'f_gen': 1017,
            'u_gen_l1': 1018, 'u_gen_l2': 1019, 'u_gen_l3': 1020,
            'i_l1': 1024, 'i_l2': 1025, 'i_l3': 1026,
            'f_mains': 1029, 'u_mains_l1': 1030, 'u_mains_l2': 1031, 'u_mains_l3': 1032,
            'battery': 1036, 'oil': 1038, 'temp': 1039, 'fuel': 1040,
            'engine_state': 1286,
        }

    async def get_snapshot(self, db, notif_id):
        try:
            cursor = db.execute("SELECT ts, controller_id FROM events WHERE id = ?", (notif_id,))
            row = cursor.fetchone()
            if not row:
                return {"error": "Notification not found"}

            notif_ts = row['ts']
            snapshot = {}

            for addr, name, unit in self.snapshot_params:
                cursor = db.execute("""
                    SELECT s.value_text FROM samples s
                    JOIN register_catalog rc ON rc.id = s.catalog_id
                    WHERE rc.addr = ? AND s.ts <= ?
                    ORDER BY s.ts DESC LIMIT 1
                """, (addr, notif_ts))
                sample = cursor.fetchone()
                if sample and sample[0] is not None:
                    try:
                        value = float(sample[0])
                        if value == 32768 or value < 0 or value > 1e6:
                            formatted = "—"
                        else:
                            if addr in (1039, 1040):
                                formatted = f"{value:.0f} {unit}".strip()
                            elif addr == 1038:
                                formatted = f"{value:.1f} {unit}".strip()
                            else:
                                formatted = f"{value:.1f} {unit}".strip()
                    except (ValueError, TypeError):
                        formatted = str(sample[0])
                else:
                    formatted = "—"
                snapshot[name] = formatted

            return snapshot
        except Exception as e:
            return {"error": str(e)}

    async def get_full_status(self, db, controller_id):
        try:
            # Получаем имя установки
            cursor = db.execute("""
                SELECT last_value_text FROM last_value lv
                JOIN register_catalog rc ON rc.id = lv.catalog_id
                WHERE rc.addr = 3000 LIMIT 1
            """)
            row = cursor.fetchone()
            unit_name = row[0] if row and row[0] else "ДГУ"

            # Получаем параметры
            params = {}
            for name, addr in self.full_status_params.items():
                cursor = db.execute("""
                    SELECT lv.last_value_text
                    FROM last_value lv
                    JOIN register_catalog rc ON rc.id = lv.catalog_id
                    WHERE rc.addr = ? LIMIT 1
                """, (addr,))
                row = cursor.fetchone()
                if row and row[0] is not None:
                    try:
                        params[name] = float(row[0])
                    except ValueError:
                        params[name] = row[0]
                else:
                    params[name] = None

            # Режимы
            mode_addrs = {54: 'mode_off', 55: 'mode_man', 56: 'mode_auto', 57: 'mode_test'}
            modes = {}
            for addr, name in mode_addrs.items():
                cursor = db.execute("""
                    SELECT lv.last_value_text
                    FROM last_value lv
                    JOIN register_catalog rc ON rc.id = lv.catalog_id
                    WHERE rc.addr = ? LIMIT 1
                """, (addr,))
                row = cursor.fetchone()
                modes[name] = int(row[0]) if row and row[0] else 0

            if modes['mode_off'] == 1:
                mode = "ВЫКЛ"
            elif modes['mode_man'] == 1:
                mode = "РУЧН"
            elif modes['mode_auto'] == 1:
                mode = "АВТО"
            elif modes['mode_test'] == 1:
                mode = "ТЕСТ"
            else:
                mode = "неизвестно"

            # Сеть
            cursor = db.execute("""
                SELECT lv.last_value_text
                FROM last_value lv
                JOIN register_catalog rc ON rc.id = lv.catalog_id
                WHERE rc.addr = 44 LIMIT 1
            """)
            row = cursor.fetchone()
            mains_ok = int(row[0]) if row and row[0] else None
            mains_state = "нет данных" if mains_ok is None else ("✅ в норме" if mains_ok == 1 else "🚨 АВАРИЯ")

            # Контакторы
            cursor = db.execute("""
                SELECT rc.addr, lv.last_value_text
                FROM last_value lv
                JOIN register_catalog rc ON rc.id = lv.catalog_id
                WHERE rc.addr IN (149, 150)
            """)
            contacts = {row[0]: row[1] for row in cursor.fetchall()}

            # Состояния двигателя
            engine_states = {
                0: "Инициализ", 1: "Готов", 2: "Не готов", 3: "Престарт",
                4: "Прокрутка", 5: "Пауза", 6: "Запуск", 7: "В работе",
                8: "Нагрузка", 9: "ПлавнРазгр", 10: "Охлаждение", 11: "Стоп",
                12: "Остановка", 13: "Продув", 14: "Экстр РУЧН", 15: "ПлавнНагр",
                16: "ОжидСтоп", 17: "ОстПродув", 18: "Возбуждение", 19: "Прогрев",
                20: "Без нагр"
            }
            engine_state_val = params.get('engine_state')
            engine_state = engine_states.get(int(engine_state_val), "—") if engine_state_val is not None else "—"

            status_data = {
                "unit_name": unit_name,
                "general": {
                    "engine_state": engine_state,
                    "mode": mode,
                    "mains_state": mains_state,
                    "mains_ok": mains_ok,
                    "contacts_avs": "ВКЛ" if contacts.get(149) == '1' else "ВЫКЛ",
                    "contacts_avg": "ВКЛ" if contacts.get(150) == '1' else "ВЫКЛ",
                },
                "engine": {
                    "rpm": self._format_value(params.get('rpm'), 0),
                    "p_load": self._format_value(params.get('p_load'), 1),
                },
                "generator": {
                    "f_gen": self._format_value(params.get('f_gen'), 1),
                    "u_gen_l1": self._format_value(params.get('u_gen_l1'), 1),
                    "u_gen_l2": self._format_value(params.get('u_gen_l2'), 1),
                    "u_gen_l3": self._format_value(params.get('u_gen_l3'), 1),
                },
                "current": {
                    "i_l1": self._format_value(params.get('i_l1'), 1),
                    "i_l2": self._format_value(params.get('i_l2'), 1),
                    "i_l3": self._format_value(params.get('i_l3'), 1),
                },
                "mains": {
                    "f_mains": self._format_value(params.get('f_mains'), 1),
                    "u_mains_l1": self._format_value(params.get('u_mains_l1'), 1),
                    "u_mains_l2": self._format_value(params.get('u_mains_l2'), 1),
                    "u_mains_l3": self._format_value(params.get('u_mains_l3'), 1),
                },
                "sensors": {
                    "battery": self._format_value(params.get('battery'), 1),
                    "oil": self._format_value(params.get('oil'), 1),
                    "temp": self._format_value(params.get('temp'), 0),
                    "fuel": self._format_value(params.get('fuel'), 0),
                }
            }

            return {
                "status_text": "",
                "status_data": status_data,
                "params": params,
                "mode": mode,
                "engine_state": engine_state,
                "mains_state": mains_state,
                "unit_name": unit_name
            }
        except Exception as e:
            return {"error": str(e)}

    async def get_bits_status(self, db, controller_id):
        try:
            important_addrs = [36, 37, 38, 39, 40, 43, 44, 69, 70, 71, 149, 150]
            placeholders = ','.join('?' * len(important_addrs))

            query = f"""
                SELECT rc.addr, rc.name, lv.last_value_text
                FROM last_value lv
                JOIN register_catalog rc ON rc.id = lv.catalog_id
                WHERE rc.kind = 'binary_bit' 
                AND rc.addr IN ({placeholders})
            """

            cursor = db.execute(query, important_addrs)
            rows = cursor.fetchall()

            result = {}
            for row in rows:
                addr = row['addr']
                raw_name = row['name'] or f"BIT_{addr}"
                state = "ВКЛ" if row['last_value_text'] == '1' else "ВЫКЛ"
                result[addr] = {
                    "name": raw_name,
                    "state": state,
                    "state_num": int(row['last_value_text'] or 0)
                }

            return {
                "status": "ok",
                "bits": result,
                "timestamp": int(datetime.now().timestamp())
            }
        except Exception as e:
            return {"error": str(e)}

    async def get_current_alerts(self, db, controller_id):
        try:
            alerts = []
            # Получаем список контроллеров
            if controller_id:
                cursor = db.execute("SELECT id, name FROM controllers WHERE id = ?", (controller_id,))
                controllers = cursor.fetchall()
            else:
                cursor = db.execute("SELECT id, name FROM controllers")
                controllers = cursor.fetchall()

            for ctrl in controllers:
                ctrl_id = ctrl["id"]
                ctrl_name = ctrl["name"]

                # Активные битовые сигналы
                cursor = db.execute("""
                    SELECT rc.name, rc.addr, lv.last_value_text
                    FROM last_value lv
                    JOIN register_catalog rc ON rc.id = lv.catalog_id
                    WHERE rc.kind='binary_bit' AND lv.last_value_text='1'
                    ORDER BY rc.addr
                """)

                for row in cursor.fetchall():
                    name, addr, val = row
                    cursor_check = db.execute("""
                        SELECT 1 FROM events 
                        WHERE code = ? 
                        LIMIT 1
                    """, (f"BIT_{addr}",))
                    if not cursor_check.fetchone():
                        continue

                    clean_name = name
                    if clean_name:
                        import re
                        clean_name = re.sub(r'^\d+\s+', '', str(name))
                        clean_name = re.sub(r'\s+', ' ', clean_name).strip()

                    title_lower = (name or '').lower()

                    if any(k in title_lower for k in ['авар', 'критич', 'alarm', 'ост', 'shutdown', 'стоп', 'рво', 'пожар']):
                        alerts.append({
                            'severity': 'alarm',
                            'code': f'BIT_{addr}',
                            'title': clean_name,
                            'details': 'Активна аварийная защита',
                            'controller_id': ctrl_id,
                            'controller_name': ctrl_name
                        })
                    elif any(k in title_lower for k in ['прд', 'пред', 'warn', 'низк', 'высок']):
                        alerts.append({
                            'severity': 'warn',
                            'code': f'BIT_{addr}',
                            'title': clean_name,
                            'details': 'Активно предупреждение',
                            'controller_id': ctrl_id,
                            'controller_name': ctrl_name
                        })

                # Критические аналоговые параметры
                cursor = db.execute("""
                    SELECT rc.addr, rc.name, lv.last_value_text
                    FROM last_value lv
                    JOIN register_catalog rc ON rc.id = lv.catalog_id
                    WHERE rc.addr IN (1036, 1038, 1039, 1040, 1029, 1017, 1030)
                    AND lv.last_value_text IS NOT NULL
                """)

                for row in cursor.fetchall():
                    addr, name, value = row
                    try:
                        val = float(value)

                        # Температура ОЖ
                        if addr == 1039 and val > 90:
                            alerts.append({
                                'severity': 'alarm' if val > 95 else 'warn',
                                'code': 'HIGH_TEMP_ALARM' if val > 95 else 'HIGH_TEMP_WARN',
                                'title': 'КРИТИЧЕСКАЯ ТЕМПЕРАТУРА ОЖ' if val > 95 else 'Высокая температура ОЖ',
                                'details': f'Температура: {val:.1f}°C',
                                'controller_id': ctrl_id,
                                'controller_name': ctrl_name
                            })

                        # Давление масла
                        elif addr == 1038 and val < 2.0:
                            alerts.append({
                                'severity': 'alarm' if val < 1.0 else 'warn',
                                'code': 'LOW_OIL_ALARM' if val < 1.0 else 'LOW_OIL_WARN',
                                'title': 'АВАРИЙНО НИЗКОЕ ДАВЛЕНИЕ МАСЛА' if val < 1.0 else 'Низкое давление масла',
                                'details': f'Давление: {val:.1f} Бар',
                                'controller_id': ctrl_id,
                                'controller_name': ctrl_name
                            })

                        # Напряжение АКБ
                        elif addr == 1036 and (val < 18 or val > 30):
                            alerts.append({
                                'severity': 'alarm' if (val < 16 or val > 32) else 'warn',
                                'code': 'LOW_BATTERY_ALARM' if val < 16 else (
                                    'HIGH_BATTERY_ALARM' if val > 32 else 'BATTERY_WARN'),
                                'title': 'КРИТИЧЕСКИ НИЗКОЕ НАПРЯЖЕНИЕ АКБ' if val < 16 else (
                                    'КРИТИЧЕСКИ ВЫСОКОЕ НАПРЯЖЕНИЕ АКБ' if val > 32 else 'Отклонение напряжения АКБ'),
                                'details': f'Напряжение: {val:.1f}В',
                                'controller_id': ctrl_id,
                                'controller_name': ctrl_name
                            })

                        # Уровень топлива
                        elif addr == 1040 and val < 20:
                            alerts.append({
                                'severity': 'alarm' if val <= 10 else 'warn',
                                'code': 'LOW_FUEL_ALARM' if val <= 10 else 'LOW_FUEL_WARN',
                                'title': 'КРИТИЧЕСКИ НИЗКИЙ УРОВЕНЬ ТОПЛИВА' if val <= 10 else 'Низкий уровень топлива',
                                'details': f'Уровень топлива: {val:.0f}%',
                                'controller_id': ctrl_id,
                                'controller_name': ctrl_name
                            })

                        # Частота
                        elif addr == 1029 and (val < 47 or val > 53):
                            alerts.append({
                                'severity': 'alarm' if (val < 47 or val > 53) else 'warn',
                                'code': 'FREQ_ALARM',
                                'title': 'КРИТИЧЕСКОЕ ОТКЛОНЕНИЕ ЧАСТОТЫ СЕТИ',
                                'details': f'Частота: {val:.1f}Гц',
                                'controller_id': ctrl_id,
                                'controller_name': ctrl_name
                            })

                        # Частота генератора (адрес 1017) - ДОБАВЛЕНО
                        elif addr == 1017 and (val < 47 or val > 53):
                            alerts.append({
                                'severity': 'alarm' if (val < 47 or val > 53) else 'warn',
                                'code': 'FREQ_GEN_ALARM',
                                'title': 'КРИТИЧЕСКОЕ ОТКЛОНЕНИЕ ЧАСТОТЫ ГЕНЕРАТОРА',
                                'details': f'Частота генератора: {val:.1f}Гц (норма 47-53Гц)',
                                'controller_id': ctrl_id,
                                'controller_name': ctrl_name
                            })

                        # Напряжение сети L1-N (адрес 1030)
                        elif addr == 1030 and (val < 207 or val > 253):
                            alerts.append({
                                'severity': 'alarm' if (val < 180 or val > 280) else 'warn',
                                'code': 'VOLTAGE_ALARM',
                                'title': 'КРИТИЧЕСКОЕ ОТКЛОНЕНИЕ НАПРЯЖЕНИЯ СЕТИ',
                                'details': f'Напряжение сети L1-N: {val:.1f}В (норма 207-253В)',
                                'controller_id': ctrl_id,
                                'controller_name': ctrl_name
                            })

                    except (ValueError, TypeError):
                        pass

            return alerts
        except Exception as e:
            logger.error(f"Ошибка получения всех аварий: {e}")
            return {"error": str(e)}

    async def export_report(self, db, controller_id, hours):
        """
        Формирует полный отчёт по установке в Excel.
        Возвращает словарь с content, media_type, filename.
        """
        try:
            wb = Workbook()

            # Получаем имя установки
            cursor = db.execute("SELECT name FROM controllers WHERE id = ?", (controller_id,))
            controller = cursor.fetchone()
            controller_name = controller['name'] if controller else f"Установка {controller_id}"

            # ----- ЛИСТ 1: Журнал событий -----
            ws_events = wb.active
            ws_events.title = "Журнал событий"
            start_ts = int((datetime.now() - timedelta(hours=hours)).timestamp())

            cursor = db.execute("""
                SELECT e.ts, e.severity, e.code, e.title, e.details
                FROM events e
                WHERE e.controller_id = ? AND e.ts > ?
                ORDER BY e.ts DESC
            """, (controller_id, start_ts))
            rows = cursor.fetchall()

            headers = ["Дата и время", "Важность", "Код", "Событие", "Детали"]
            for col, header in enumerate(headers, 1):
                cell = ws_events.cell(row=1, column=col, value=header)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
                cell.alignment = Alignment(horizontal="center")

            for row_idx, row in enumerate(rows, 2):
                dt = datetime.fromtimestamp(row['ts']) + timedelta(hours=3)
                ws_events.cell(row=row_idx, column=1, value=dt.strftime("%d.%m.%Y %H:%M:%S"))

                severity_cell = ws_events.cell(row=row_idx, column=2, value=row['severity'].upper())
                if row['severity'] == 'alarm':
                    severity_cell.fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
                    severity_cell.font = Font(color="FFFFFF", bold=True)
                elif row['severity'] == 'warn':
                    severity_cell.fill = PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid")
                    severity_cell.font = Font(color="FFFFFF", bold=True)

                ws_events.cell(row=row_idx, column=3, value=row['code'] or '')
                ws_events.cell(row=row_idx, column=4, value=row['title'])
                ws_events.cell(row=row_idx, column=5, value=row['details'] or '')

            # Автоширина для событий
            for column in ws_events.columns:
                max_length = 0
                col_letter = column[0].column_letter
                for cell in column:
                    try:
                        length = len(str(cell.value))
                        if length > max_length:
                            max_length = length
                    except:
                        pass
                ws_events.column_dimensions[col_letter].width = min(max_length + 2, 50)

            # ----- ЛИСТ 2: Данные графиков -----
            ws_charts = wb.create_sheet("Данные графиков")
            parameters = [
                (1000, "Обороты (RPM)", "об/мин"),
                (1039, "Температура ОЖ", "°C"),
                (1038, "Давление масла", "Бар"),
                (1040, "Уровень топлива", "%"),
                (1036, "Напряжение АКБ", "В"),
                (1017, "Частота генератора", "Гц"),
            ]

            ws_charts.cell(row=1, column=1, value="Дата и время")
            for col, (_, name, unit) in enumerate(parameters, 2):
                ws_charts.cell(row=1, column=col, value=f"{name}, {unit}")
                ws_charts.cell(row=1, column=col).font = Font(bold=True)

            all_data = {}
            timestamps = set()

            for idx, (addr, _, _) in enumerate(parameters):
                cursor = db.execute("""
                    SELECT s.ts, s.value_text
                    FROM samples s
                    JOIN register_catalog rc ON rc.id = s.catalog_id
                    WHERE rc.addr = ? AND s.ts > ?
                    ORDER BY s.ts
                """, (addr, start_ts))

                rows = cursor.fetchall()
                for row in rows:
                    dt = datetime.fromtimestamp(row['ts']) + timedelta(hours=3)
                    ts_str = dt.strftime("%d.%m.%Y %H:%M:%S")
                    timestamps.add(ts_str)
                    try:
                        value = float(row[1]) if row[1] else None
                        all_data[(ts_str, idx)] = f"{value:.1f}" if value is not None else "—"
                    except:
                        all_data[(ts_str, idx)] = "—"

            sorted_timestamps = sorted(timestamps)
            for row_idx, ts in enumerate(sorted_timestamps, 2):
                ws_charts.cell(row=row_idx, column=1, value=ts)
                for col_idx in range(len(parameters)):
                    value = all_data.get((ts, col_idx), "—")
                    ws_charts.cell(row=row_idx, column=col_idx + 2, value=value)

            # ----- ЛИСТ 3: Текущий статус -----
            ws_status = wb.create_sheet("Текущий статус")
            ws_status.cell(row=1, column=1, value="Параметр")
            ws_status.cell(row=1, column=2, value="Значение")
            ws_status.cell(row=1, column=3, value="Единица")
            ws_status.cell(row=1, column=1).font = Font(bold=True)
            ws_status.cell(row=1, column=2).font = Font(bold=True)
            ws_status.cell(row=1, column=3).font = Font(bold=True)

            status_params = [
                (1000, "Обороты", "об/мин"),
                (1039, "Температура ОЖ", "°C"),
                (1038, "Давление масла", "Бар"),
                (1040, "Уровень топлива", "%"),
                (1036, "Напряжение АКБ", "В"),
                (1029, "Частота сети", "Гц"),
                (1017, "Частота генератора", "Гц"),
                (1018, "Напряжение L1-N", "В"),
                (1019, "Напряжение L2-N", "В"),
                (1020, "Напряжение L3-N", "В"),
            ]

            row = 2
            for addr, name, unit in status_params:
                cursor = db.execute("""
                    SELECT lv.last_value_text
                    FROM last_value lv
                    JOIN register_catalog rc ON rc.id = lv.catalog_id
                    WHERE rc.addr = ?
                """, (addr,))
                result = cursor.fetchone()
                value = result[0] if result and result[0] else "—"
                ws_status.cell(row=row, column=1, value=name)
                ws_status.cell(row=row, column=2, value=value)
                ws_status.cell(row=row, column=3, value=unit)
                row += 1

            # ----- ЛИСТ 4: Статистика -----
            ws_stats = wb.create_sheet("Статистика")
            cursor = db.execute("""
                SELECT 
                    COUNT(*) as total,
                    SUM(CASE WHEN severity = 'alarm' THEN 1 ELSE 0 END) as alarms,
                    SUM(CASE WHEN severity = 'warn' THEN 1 ELSE 0 END) as warns
                FROM events
                WHERE controller_id = ? AND ts > ?
            """, (controller_id, start_ts))
            stats = cursor.fetchone()

            ws_stats.cell(row=1, column=1, value="Показатель")
            ws_stats.cell(row=1, column=2, value="Значение")
            ws_stats.cell(row=1, column=1).font = Font(bold=True)
            ws_stats.cell(row=1, column=2).font = Font(bold=True)

            ws_stats.cell(row=2, column=1, value="Всего событий")
            ws_stats.cell(row=2, column=2, value=stats['total'] or 0)
            ws_stats.cell(row=3, column=1, value="Аварий (alarm)")
            ws_stats.cell(row=3, column=2, value=stats['alarms'] or 0)
            ws_stats.cell(row=4, column=1, value="Предупреждений (warn)")
            ws_stats.cell(row=4, column=2, value=stats['warns'] or 0)

            ws_stats.cell(row=6, column=1, value="Период отчета")
            ws_stats.cell(row=6, column=2, value=f"Последние {hours} часов")
            ws_stats.cell(row=7, column=1, value="Дата формирования")
            ws_stats.cell(row=7, column=2, value=datetime.now().strftime("%d.%m.%Y %H:%M:%S"))

            # Сохраняем в BytesIO
            output = io.BytesIO()
            wb.save(output)
            output.seek(0)

            filename = f"report_{controller_name}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            return {
                "content": output.getvalue(),
                "media_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "filename": filename
            }

        except Exception as e:
            logger.error(f"Ошибка создания полного отчета: {e}")
            return {"error": str(e)}

    def _format_value(self, value, decimals):
        if value is None:
            return "—"
        return f"{value:.{decimals}f}"