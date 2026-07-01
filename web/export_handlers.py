"""
Обработчики экспорта данных для веб-интерфейса
"""

import csv
import io
from typing import List, Dict
from datetime import datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
import sqlite3


def export_to_csv_buffer(data: List[Dict]) -> io.StringIO:
    """
    Создаёт CSV-строку в StringIO.
    """
    if not data:
        data = [{"Сообщение": "Нет данных для экспорта"}]

    output = io.StringIO()
    fieldnames = data[0].keys() if data else []
    writer = csv.DictWriter(output, fieldnames=fieldnames, delimiter=';')
    writer.writeheader()
    writer.writerows(data)
    output.seek(0)
    return output


def export_to_excel_buffer(data: List[Dict]) -> io.BytesIO:
    """
    Создаёт Excel-файл в BytesIO.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Экспорт"

    if data:
        headers = list(data[0].keys())
        # Заголовки
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")

        # Данные
        for row_idx, row in enumerate(data, 2):
            for col_idx, header in enumerate(headers, 1):
                value = row.get(header, "")
                ws.cell(row=row_idx, column=col_idx, value=value)

        # Автоширина
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    length = len(str(cell.value))
                    if length > max_length:
                        max_length = length
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def prepare_events_data(rows: List[sqlite3.Row]) -> List[Dict]:
    """
    Преобразует строки из БД в список словарей для экспорта.
    (Полностью скопировано из вашего рабочего кода)
    """
    export_data = []
    for row in rows:
        dt = datetime.fromtimestamp(row['ts']) + timedelta(hours=3)

        # Формируем имя установки
        if row['modbus_name'] and row['modbus_name'].strip():
            controller_display = row['modbus_name']
        elif row['name'] and row['name'].strip():
            controller_display = row['name']
        elif row['ip'] and row['ip'].strip():
            controller_display = f"ДГУ ({row['ip']})"
        else:
            controller_display = f"ДГУ {row.get('controller_id', '?')}"

        export_data.append({
            'Дата и время': dt.strftime("%d.%m.%Y %H:%M:%S"),
            'Установка': controller_display,
            'Важность': row['severity'].upper() if row['severity'] else '',
            'Код': row['code'] or '',
            'Событие': row['title'],
            'Детали': row['details'] or ''
        })
    return export_data