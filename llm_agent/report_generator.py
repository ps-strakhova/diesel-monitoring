"""
Генератор актов строго по шаблонам компании
"""

import io
from datetime import datetime
from typing import Dict, List, Optional
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


class ReportGenerator:
    """Генерация актов по шаблонам компании"""

    def __init__(self, db_path: str = None):
        self.db_path = db_path

    def get_company_header(self, doc):
        """Добавление шапки компании в документ"""

        # Шапка с логотипом и названием компании
        header_table = doc.add_table(rows=2, cols=2)
        header_table.columns[0].width = Cm(3)
        header_table.columns[1].width = Cm(13)

        # Левая ячейка (для логотипа)
        left_cell = header_table.cell(0, 0)
        left_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        # Правая ячейка (название компании)
        right_cell = header_table.cell(0, 1)
        right_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        company_name = right_cell.paragraphs[0]
        company_name.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = company_name.add_run("ООО «Компания Дизель»")
        run.bold = True
        run.font.size = Pt(14)

        # Контактная информация
        contact_para = right_cell.add_paragraph()
        contact_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        contact_para.add_run("150044, г. Ярославль, пр-кт Ленинградский, зд. 33в, помещ. 401\n")
        contact_para.add_run("Тел.: +7 (800) 5112-461, E-mail: info@comd.ru")
        contact_para.runs[0].font.size = Pt(8)
        contact_para.runs[1].font.size = Pt(8)

        # Вторая строка шапки
        header_table.cell(1, 0).merge(header_table.cell(1, 1))
        bottom_cell = header_table.cell(1, 0)
        bottom_para = bottom_cell.paragraphs[0]
        bottom_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        bottom_para.add_run("ДИЗЕЛЬНЫЕ ЭЛЕКТРОСТАНЦИИ • СЕРВИСНОЕ ОБСЛУЖИВАНИЕ • РЕМОНТ")
        bottom_para.runs[0].font.size = Pt(8)
        bottom_para.runs[0].italic = True

        doc.add_paragraph()

        return header_table

    def generate_claim_act(self, monitoring_data: Dict,
                           conversation_history: List[Dict],
                           user_info: Dict = None) -> io.BytesIO:
        """Генерация Акта-Заявки по неисправности (строго по шаблону)"""

        doc = Document()

        # Устанавливаем поля страницы
        section = doc.sections[0]
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(2)
        section.right_margin = Cm(1.5)

        # ШАПКА КОМПАНИИ
        self.get_company_header(doc)

        # ЗАГОЛОВОК АКТА
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title.add_run("Акт-Заявка по неисправности")
        title_run.bold = True
        title_run.font.size = Pt(16)
        title_run.underline = True

        # ДАТА
        date_para = doc.add_paragraph()
        date_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        date_para.add_run(f'«{datetime.now().strftime("%d")}» ____________ 20__ г.')

        doc.add_paragraph()

        # ========== КОНТАКТНАЯ ИНФОРМАЦИЯ ==========
        doc.add_heading("Контактная информация", level=2)

        contact_table = doc.add_table(rows=5, cols=2)
        contact_table.style = 'Table Grid'
        contact_table.alignment = WD_TABLE_ALIGNMENT.CENTER

        # Ширина колонок
        contact_table.columns[0].width = Cm(5)
        contact_table.columns[1].width = Cm(11)

        contact_data = [
            ("Заявитель, адрес:",
             user_info.get('applicant', '________________________') if user_info else '________________________'),
            ("ИНН:", user_info.get('inn', '________________________') if user_info else '________________________'),
            ("Потребитель, адрес:", user_info.get('consumer', monitoring_data.get('unit_name',
                                                                                  '________________________')) if user_info else monitoring_data.get(
                'unit_name', '________________________')),
            ("Контактное лицо:",
             user_info.get('contact_person', '________________________') if user_info else '________________________'),
            ("Телефон, эл. почта:",
             user_info.get('phone', '________________________') if user_info else '________________________'),
        ]

        for i, (label, value) in enumerate(contact_data):
            contact_table.cell(i, 0).text = label
            contact_table.cell(i, 0).paragraphs[0].runs[0].bold = True
            contact_table.cell(i, 1).text = str(value)

        doc.add_paragraph()

        # ========== ИНФОРМАЦИЯ ОБ ОБОРУДОВАНИИ ==========
        doc.add_heading("Информация об оборудовании", level=2)

        equip_table = doc.add_table(rows=6, cols=2)
        equip_table.style = 'Table Grid'
        equip_table.columns[0].width = Cm(5)
        equip_table.columns[1].width = Cm(11)

        # Получаем наработку и другие данные
        hours = monitoring_data.get('hours', monitoring_data.get('p_load', '—'))
        last_to = monitoring_data.get('last_maintenance', '—')

        equip_data = [
            ("Модель оборудования:", monitoring_data.get('model', monitoring_data.get('unit_name', 'ДГУ'))),
            ("Серийный №:", monitoring_data.get('serial', '________________________')),
            ("Дата выпуска (число/месяц/год):",
             monitoring_data.get('manufacture_date', '________ / ________ / ________')),
            ("Дата ввода в эксплуатацию:", monitoring_data.get('commissioning_date', '________ / ________ / ________')),
            ("Наработка (м/ч):", f"{hours}" if hours != '—' else '________'),
            ("Последнее ТО, дата:", last_to if last_to != '—' else '________ / ________ / ________'),
        ]

        for i, (label, value) in enumerate(equip_data):
            equip_table.cell(i, 0).text = label
            equip_table.cell(i, 0).paragraphs[0].runs[0].bold = True
            equip_table.cell(i, 1).text = str(value)

        doc.add_paragraph()

        # ========== ХАРАКТЕР НЕИСПРАВНОСТИ ==========
        doc.add_heading("Характер неисправности, внешние проявления, обстоятельства появления*", level=2)

        # Извлекаем проблему из диалога
        problem_text = self._extract_problem_from_dialogue(conversation_history)

        problem_para = doc.add_paragraph()
        problem_para.add_run(problem_text)

        # Текущие параметры установки
        doc.add_paragraph()
        doc.add_heading("Текущие параметры установки:", level=3)

        params_text = f"""
- Обороты: {monitoring_data.get('rpm', '—')} об/мин
- Температура ОЖ: {monitoring_data.get('temp', '—')}°C
- Давление масла: {monitoring_data.get('oil', '—')} Бар
- Напряжение АКБ: {monitoring_data.get('battery', '—')} В
- Уровень топлива: {monitoring_data.get('fuel', '—')} %
- Режим работы: {'АВАРИЙНЫЙ' if monitoring_data.get('alarm_mode') else 'НОРМАЛЬНЫЙ'}
"""
        doc.add_paragraph(params_text)

        doc.add_paragraph()

        # ========== ПРЕДУПРЕЖДЕНИЕ ==========
        warning = doc.add_paragraph()
        warning_run = warning.add_run(
            "ВНИМАНИЕ! В случае признания неисправности не гарантийным случаем, Заявитель обязуется возместить все расходы, связанные с выездом представителей ООО «Компания Дизель» и рассмотрением рекламации по данному обращению.")
        warning_run.italic = True
        warning_run.font.size = Pt(9)

        doc.add_paragraph()

        # ========== ПОДПИСИ ==========
        sign_table = doc.add_table(rows=2, cols=2)
        sign_table.style = 'Table Grid'
        sign_table.columns[0].width = Cm(8)
        sign_table.columns[1].width = Cm(8)

        sign_table.cell(0, 0).text = "Ф.И.О. и должность лица, заполнившего заявку:"
        sign_table.cell(0, 0).paragraphs[0].runs[0].bold = True
        sign_table.cell(0, 1).text = user_info.get('contact_person',
                                                   '________________________') if user_info else '________________________'

        sign_table.cell(1, 0).text = "М.П."
        sign_table.cell(1, 0).paragraphs[0].runs[0].bold = True
        sign_table.cell(1, 1).text = "Подпись: _________________"

        doc.add_paragraph()

        # Примечание
        note = doc.add_paragraph()
        note_run = note.add_run(
            "*ВНИМАНИЕ! Все графы обязательны к заполнению. При невыполнении данного условия заявка к рассмотрению не принимается.")
        note_run.italic = True
        note_run.font.size = Pt(8)

        # Сохраняем
        output = io.BytesIO()
        doc.save(output)
        output.seek(0)

        return output

    def generate_work_order(self, monitoring_data: Dict,
                            work_list: List[str] = None,
                            materials: List[Dict] = None) -> io.BytesIO:
        """Генерация Заказ-наряда"""

        doc = Document()

        # Устанавливаем поля
        section = doc.sections[0]
        section.top_margin = Cm(1.5)
        section.left_margin = Cm(1.5)

        # Шапка
        self.get_company_header(doc)

        # Заголовок
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc_number = datetime.now().strftime("%Y%m%d")
        title_run = title.add_run(f"Заказ-наряд № ТО_{doc_number}")
        title_run.bold = True
        title_run.font.size = Pt(14)
        title_run.underline = True

        doc.add_paragraph()

        # Информация о заказчике
        info_table = doc.add_table(rows=6, cols=2)
        info_table.style = 'Table Grid'
        info_table.columns[0].width = Cm(5)
        info_table.columns[1].width = Cm(11)

        hours = monitoring_data.get('hours', monitoring_data.get('p_load', '—'))

        info_data = [
            ("Заказчик:", monitoring_data.get('customer', '________________________')),
            ("Адрес:", monitoring_data.get('address', '________________________')),
            ("Потребитель:",
             monitoring_data.get('consumer', monitoring_data.get('unit_name', '________________________'))),
            ("Модель оборудования:", monitoring_data.get('model', monitoring_data.get('unit_name', 'ДГУ'))),
            ("Сер. № ДВС / Генератора:",
             f"{monitoring_data.get('serial', '—')} / {monitoring_data.get('gen_serial', '—')}"),
            ("Наработка:", f"{hours} м/ч" if hours != '—' else '________ м/ч'),
        ]

        for i, (label, value) in enumerate(info_data):
            info_table.cell(i, 0).text = label
            info_table.cell(i, 0).paragraphs[0].runs[0].bold = True
            info_table.cell(i, 1).text = str(value)

        doc.add_paragraph()

        # Выполненные работы
        doc.add_heading("Выполненные операции (работы)", level=2)

        work_table = doc.add_table(rows=1 + len(work_list or []), cols=3)
        work_table.style = 'Table Grid'

        # Заголовки
        headers = ["№", "Выполненные операции (работы)", "Выполнение"]
        for i, header in enumerate(headers):
            cell = work_table.cell(0, i)
            cell.text = header
            cell.paragraphs[0].runs[0].bold = True

        if work_list:
            for i, work in enumerate(work_list, 1):
                work_table.cell(i, 0).text = str(i)
                work_table.cell(i, 1).text = work
                work_table.cell(i, 2).text = "✅"
        else:
            # Пустые строки для заполнения
            for i in range(1, 6):
                work_table.cell(i, 0).text = str(i)
                work_table.cell(i, 1).text = "________________________"
                work_table.cell(i, 2).text = "□"

        doc.add_paragraph()

        # Использованные материалы
        doc.add_heading("Использованные материалы", level=2)

        mat_table = doc.add_table(rows=1 + len(materials or []), cols=4)
        mat_table.style = 'Table Grid'

        mat_headers = ["№", "Наименование материалов", "Ед.", "Использовано"]
        for i, header in enumerate(mat_headers):
            cell = mat_table.cell(0, i)
            cell.text = header
            cell.paragraphs[0].runs[0].bold = True

        if materials:
            for i, mat in enumerate(materials, 1):
                mat_table.cell(i, 0).text = str(i)
                mat_table.cell(i, 1).text = mat.get('name', '________________________')
                mat_table.cell(i, 2).text = mat.get('unit', 'шт')
                mat_table.cell(i, 3).text = str(mat.get('quantity', '____'))
        else:
            for i in range(1, 5):
                mat_table.cell(i, 0).text = str(i)
                mat_table.cell(i, 1).text = "________________________"
                mat_table.cell(i, 2).text = "____"
                mat_table.cell(i, 3).text = "____"

        doc.add_paragraph()

        # Рекомендации
        doc.add_heading("Рекомендации (примечания)", level=2)
        doc.add_paragraph("_________________________________________________________________")
        doc.add_paragraph("_________________________________________________________________")

        doc.add_paragraph()

        # Даты и подписи
        date_para = doc.add_paragraph()
        date_para.add_run(f"Работы проводились в период с \"__\" __________ 20__ г. по \"__\" __________ 20__ г.")

        doc.add_paragraph()
        doc.add_paragraph(
            "Работы выполнены в полном объеме, потребитель не имеет претензий к качеству выполненных работ.")

        sign_para = doc.add_paragraph()
        sign_para.add_run(
            f"Дата сдачи работ: \"{datetime.now().strftime('%d')}\" __________ {datetime.now().strftime('%Y')} г.")

        doc.add_paragraph()
        sign_table = doc.add_table(rows=3, cols=3)
        sign_table.cell(0, 0).text = "Должность"
        sign_table.cell(0, 1).text = "Ф.И.О."
        sign_table.cell(0, 2).text = "Подпись"
        for i in range(1, 3):
            sign_table.cell(i, 0).text = "________________________"
            sign_table.cell(i, 1).text = "________________________"
            sign_table.cell(i, 2).text = "________________________"

        output = io.BytesIO()
        doc.save(output)
        output.seek(0)

        return output

    def _extract_problem_from_dialogue(self, conversation: List[Dict]) -> str:
        """Извлечение описания проблемы из диалога"""
        if not conversation:
            return "_________________________________________________________________\n_________________________________________________________________"

        problems = []
        for msg in conversation[-10:]:
            if msg.get('role') == 'user':
                text = msg.get('text', '')
                if len(text) > 20:
                    problems.append(text)

        if problems:
            return "\n".join(problems[:3])

        return "_________________________________________________________________\n_________________________________________________________________"

    def _month_name(self, month: int) -> str:
        months = ['января', 'февраля', 'марта', 'апреля', 'мая', 'июня',
                  'июля', 'августа', 'сентября', 'октября', 'ноября', 'декабря']
        return months[month - 1]